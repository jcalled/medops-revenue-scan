"""
Homologação real: o faturamento do hospital confere, AIH por AIH, o que o sistema
diz das rejeições de um mês — e disso sai quanto cada regra acerta.

O que se julga em cada AIH:
- FATURASUS: quando o FaturaSUS apontou o erro com o dado público, o achado da
  regra que disparou (mede a regra do FaturaSUS);
- KIT: quando não apontou, a correção que o kit indica para o motivo (mede o kit).

Mesmo método do núcleo (app/services/homologacao/confiabilidade.py): acerto só
entre quem julgou se procede; piso do intervalo de Wilson (95%) em vez da
porcentagem crua; amostra mínima de 20; corrigir sozinha só com piso ≥ 90%.
PARCIAL conta como erro — a correção não saiu inteira certa; NAO_SEI fica fora.
"""
from __future__ import annotations

import csv
import io
import math
import re
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any

from sqlalchemy.orm import Session

from app.domain.kits_motivo import kits_usados
from app.domain.prova import aih_rejeitadas
from app.domain.relatorio_recuperacao import _apontamento, _prevencao
from app.domain.resumo import _nomes
from app.models import HomologationBatch, HomologationItem

VEREDITOS = {
    "CERTO": "O sistema acertou o erro e a correção",
    "PARCIAL": "Acertou o erro, mas a correção precisa de ajuste",
    "ERRADO": "Não era esse o erro, ou a correção está errada",
    "NAO_SEI": "Não dá para conferir",
}
ACERTO = ("CERTO",)
ERRO = ("ERRADO", "PARCIAL")
AMOSTRA_MINIMA = 20
PISO_PARA_AUTOMATIZAR = 0.90


def wilson_inferior(acertos: int, total: int, z: float = 1.96) -> float:
    if total <= 0:
        return 0.0
    p = acertos / total
    denominador = 1 + z * z / total
    centro = p + z * z / (2 * total)
    margem = z * math.sqrt((p * (1 - p) + z * z / (4 * total)) / total)
    return max(0.0, (centro - margem) / denominador)


def _resumo(chave: str, contagem: dict[str, int], nome: str | None = None) -> dict[str, Any]:
    acertos = sum(contagem.get(v, 0) for v in ACERTO)
    erros = sum(contagem.get(v, 0) for v in ERRO)
    julgados = acertos + erros
    piso = wilson_inferior(acertos, julgados)
    if julgados < AMOSTRA_MINIMA:
        veredito, leitura = "SEM_DADOS", f"Ainda não dá para dizer: {julgados} parecer(es), faltam {AMOSTRA_MINIMA - julgados} para a amostra mínima."
    elif piso >= PISO_PARA_AUTOMATIZAR:
        veredito, leitura = "CONFIAVEL", f"Mesmo no pior caso plausível, acerta {piso:.0%}: candidata a corrigir sozinha."
    else:
        veredito, leitura = "EM_HOMOLOGACAO", f"No pior caso plausível acerta {piso:.0%}: segue com o faturista conferindo."
    return {"chave": chave, "nome": nome or chave, "julgados": julgados, "certos": contagem.get("CERTO", 0),
            "parciais": contagem.get("PARCIAL", 0), "errados": contagem.get("ERRADO", 0),
            "nao_sei": contagem.get("NAO_SEI", 0), "acerto": round(acertos / julgados, 4) if julgados else None,
            "piso": round(piso, 4), "veredito": veredito, "leitura": leitura}


def montar_lote(db: Session, *, cnes: list[str], competencia: str, titulo: str, organization_id: int | None,
                criado_por: int | None) -> HomologationBatch:
    """As AIH rejeitadas no processamento `competencia`, com o que o FaturaSUS disse ou o que o kit orienta."""
    linhas = aih_rejeitadas(db, cnes, [competencia])
    linhas = [l for l in linhas if l["competencia"] == competencia]
    prevencao = _prevencao(db, linhas)
    kits = kits_usados(db, {m["codigo"] for l in linhas for m in l["motivos"]})
    lote = HomologationBatch(titulo=titulo, organization_id=organization_id, cnes=sorted(set(cnes)),
                             competencia=competencia, criado_por=criado_por)
    for l in sorted(linhas, key=lambda l: (l["cnes"], -l["valor"])):
        codigos = [m["codigo"] for m in l["motivos"]]
        visto = _apontamento(prevencao.get((l["n_aih"], l["competencia"])))
        kit = next((kits[c] for c in codigos if c in kits), None)
        correcao = " ".join((kit or {}).get("passos") or [])[:1500] or "Sem orientação cadastrada para o motivo."
        if visto["grupo"] == "PEGARIA":
            origem, regras, diz = "FATURASUS", visto["regras"], visto["mensagem"] or "O FaturaSUS apontou o erro."
        else:
            descricoes = "; ".join(f"{m['codigo']} {m['descricao'] or 'sem descrição oficial'}" for m in l["motivos"])
            origem, regras, diz = "KIT", [f"KIT:{c}" for c in codigos], f"Motivo oficial: {descricoes}"
        lote.itens.append(HomologationItem(n_aih=l["n_aih"], cnes=l["cnes"], valor=l["valor"], motivos=codigos,
                                           origem=origem, regras=regras, o_que_diz=diz[:3000], correcao=correcao))
    db.add(lote)
    db.commit()
    return lote


def confiabilidade(lote: HomologationBatch) -> dict[str, Any]:
    geral: dict[str, int] = defaultdict(int)
    faturasus: dict[str, int] = defaultdict(int)
    por_regra: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    por_motivo: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    respondidos = 0
    for item in lote.itens:
        if not item.veredito:
            continue
        respondidos += 1
        geral[item.veredito] += 1
        if item.origem == "FATURASUS":
            faturasus[item.veredito] += 1
            for regra in item.regras or []:
                por_regra[regra][item.veredito] += 1
        for motivo in item.motivos or []:
            por_motivo[motivo][item.veredito] += 1
    return {
        "itens": len(lote.itens), "respondidos": respondidos,
        "do_faturasus": sum(1 for i in lote.itens if i.origem == "FATURASUS"),
        "geral": _resumo("GERAL", geral, "Tudo o que o sistema disse"),
        "faturasus": _resumo("FATURASUS", faturasus, "Achados do FaturaSUS"),
        "por_regra": sorted((_resumo(r, c) for r, c in por_regra.items()), key=lambda r: (r["piso"], -r["julgados"])),
        "por_motivo": sorted((_resumo(m, c) for m, c in por_motivo.items()), key=lambda r: (r["piso"], -r["julgados"])),
    }


CABECALHO = ["AIH", "CNES", "Hospital", "Processamento", "Valor", "Motivos", "O que o sistema diz", "Correção proposta",
             "Regra", "Veredito (CERTO, PARCIAL, ERRADO, NAO_SEI)", "Comentário", "Quem conferiu"]


def planilha(db: Session, lote: HomologationBatch) -> str:
    nomes = _nomes(db, lote.cnes)
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";")
    escritor.writerow(CABECALHO)
    for i in lote.itens:
        escritor.writerow([i.n_aih, i.cnes, nomes.get(i.cnes) or "", lote.competencia, f"{float(i.valor):.2f}".replace(".", ","),
                           ", ".join(i.motivos), i.o_que_diz, i.correcao, ", ".join(i.regras), i.veredito or "",
                           i.comentario or "", i.respondido_por or ""])
    return "﻿" + saida.getvalue()


def _linhas(conteudo: bytes, nome: str) -> list[list[str]]:
    if nome.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        planilha_xlsx = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        return [["" if c is None else str(c) for c in linha] for linha in planilha_xlsx.worksheets[0].iter_rows(values_only=True)]
    texto = conteudo.decode("utf-8-sig", errors="replace") if b"\xc3" in conteudo[:20000] or conteudo[:3] == b"\xef\xbb\xbf" else conteudo.decode("latin-1")
    primeira = texto.split("\n", 1)[0]
    separador = max((";", "\t", ","), key=primeira.count)
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _normal(valor: str) -> str:
    texto = re.sub(r"[^A-Z]", "", valor.upper().replace("Ã", "A").replace("É", "E"))
    for v in VEREDITOS:
        if texto == v.replace("_", ""):
            return v
    return {"OK": "CERTO", "SIM": "CERTO", "CORRETO": "CERTO", "NAO": "ERRADO", "INCORRETO": "ERRADO",
            "NS": "NAO_SEI", "NAOSEI": "NAO_SEI"}.get(texto, "")


def importar(db: Session, lote: HomologationBatch, conteudo: bytes, nome: str) -> dict[str, Any]:
    """Aplica os vereditos da planilha preenchida, pela AIH. Veredito vazio não apaga o que já estava."""
    linhas = [l for l in _linhas(conteudo, nome) if any(str(c).strip() for c in l)]
    if not linhas:
        return {"aplicados": 0, "ignorados": [], "sem_veredito": 0}
    cab = [str(c).strip().lower() for c in linhas[0]]
    i_aih = next((i for i, c in enumerate(cab) if "aih" in c), None)
    if i_aih is None:
        raise ValueError("A planilha precisa da coluna AIH, como na planilha baixada do lote.")
    i_ver = next((i for i, c in enumerate(cab) if c.startswith("veredito")), None)
    i_com = next((i for i, c in enumerate(cab) if c.startswith("coment")), None)
    i_quem = next((i for i, c in enumerate(cab) if c.startswith("quem")), None)
    if i_ver is None:
        raise ValueError("A planilha precisa da coluna Veredito, como na planilha baixada do lote.")
    por_aih = {i.n_aih: i for i in lote.itens}
    aplicados, sem_veredito, ignorados = 0, 0, []
    agora = datetime.now(timezone.utc)
    for linha in linhas[1:]:
        pega = (lambda i: str(linha[i]).strip() if i is not None and i < len(linha) else "")
        bruto = pega(i_aih)
        n_aih = re.sub(r"\D", "", bruto)
        # Pelo texto exato; se o Excel reformatou (pontos, espaços, notação), pelos dígitos.
        item = por_aih.get(bruto) or por_aih.get(n_aih)
        n_aih = n_aih or bruto
        veredito = _normal(pega(i_ver))
        if item is None:
            if n_aih:
                ignorados.append(n_aih)
            continue
        if not veredito:
            sem_veredito += 1
            continue
        item.veredito, item.respondido_em = veredito, agora
        item.comentario = pega(i_com) or item.comentario
        item.respondido_por = pega(i_quem)[:120] or item.respondido_por
        aplicados += 1
    db.commit()
    return {"aplicados": aplicados, "sem_veredito": sem_veredito, "ignorados": ignorados[:50]}
