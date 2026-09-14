"""
Descrição dos motivos de rejeição do SIH (tabela 0027 da auxiliar TAB_SIH).

Sem ela, o scan diria "060082" onde precisa dizer "quantidade de diárias
superior à capacidade instalada".
"""
from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from pathlib import Path

from app.adapters import datasus

DIRETORIO_AUXILIAR = "/dissemin/publicos/SIHSUS/200801_/Auxiliar"
ARQUIVO = "TAB_SIH.zip"
_TABELA_MOTIVOS = "0027"


def ler_codigos_de_erro(caminho_zip: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    with zipfile.ZipFile(caminho_zip) as pacote:
        nome = next((n for n in pacote.namelist() if n.lower().endswith("erroebloqueio.xlsx")), None)
        if nome is None:
            raise datasus.ArquivoIndisponivel(f"{Path(caminho_zip).name} não traz erroebloqueio.xlsx")
        dados = io.BytesIO(pacote.read(nome))

    # Sem read_only: a planilha publicada declara a dimensão errada, e no modo
    # somente leitura o openpyxl entrega só a primeira coluna. O arquivo é pequeno.
    planilha = load_workbook(dados, read_only=False, data_only=True).worksheets[0]
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = [str(c or "").strip() for c in next(linhas)]
    i_tab, i_item, i_desc = (cabecalho.index(c) for c in ("CO_TAB", "CO_ITEM", "DS_DESCRICAO"))

    saida: dict[str, str] = {}
    for linha in linhas:
        if str(linha[i_tab] or "").strip().zfill(4) != _TABELA_MOTIVOS:
            continue
        item = linha[i_item]
        # Célula numérica perde o zero à esquerda: 60082 é o motivo 060082.
        codigo = str(int(item)).zfill(6) if isinstance(item, (int, float)) else str(item or "").strip()
        if codigo:
            saida[codigo] = " ".join(str(linha[i_desc] or "").split())
    return saida


def obter_tab_sih(destino: Path, *, pasta_local: Path | None = None,
                  baixar: Callable[..., Path] = datasus.baixar) -> Path:
    if pasta_local:
        return datasus.achar_na_pasta(Path(pasta_local), ARQUIVO)
    return baixar(DIRETORIO_AUXILIAR, ARQUIVO, Path(destino), minimo_bytes=10_000)
